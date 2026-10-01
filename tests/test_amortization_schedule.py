"""Create an amortization schedule from a loan, file it, and email the borrower."""

import os
import tempfile
import uuid
from datetime import date

os.environ.setdefault("DATA_DIR", tempfile.mkdtemp(prefix="brittco-amort-"))
os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.setdefault("PLATFORM_MODE", "brittco_existing")
os.environ.pop("SMTP_HOST", None)

from openpyxl import load_workbook

from amortization import amortization_document_name, build_amortization_schedule
import app as brittco


def _staff():
    client = brittco.app.test_client()
    with client.session_transaction() as sess:
        sess["staff_id"] = 1
    return client


def _insert_loan(name, email, **overrides):
    fields = {
        "loan_number": "BC-AMORT-" + uuid.uuid4().hex[:8],
        "original_principal": 120000,
        "current_balance": 120000,
        "rate": 12,
        "pricing_mode": "rate",
        "flat_fee": None,
        "payment_amount": 1200,
        "payment_type": "Interest only",
        "payment_frequency": "Monthly",
        "start_date": "2026-01-15",
        "maturity_date": "2026-04-15",
        "next_payment_due": "2026-02-15",
        "base_term_months": 3,
        "property_address": "10 Oak Street",
        "status": "Current",
        "loan_type": "Fix and Flip",
    }
    fields.update(overrides)
    with brittco.app.app_context():
        brittco.db().execute(
            """INSERT INTO borrowers
               (name, entity_type, entity_name, email, phone, credit_score, password, notes)
               VALUES (?,?,?,?,?,?,?,?)""",
            (name, "LLC", name + " LLC", email, "555", 700, "borrower", ""),
        )
        bid = brittco.db().execute(
            "SELECT id FROM borrowers WHERE email=?", (email,)
        ).fetchone()["id"]
        brittco.db().execute(
            """INSERT INTO loans (
                borrower_id, loan_number, loan_type, property_address,
                original_principal, current_balance, rate, pricing_mode, flat_fee,
                start_date, maturity_date, payment_type, payment_amount,
                payment_frequency, next_payment_due, status, base_term_months
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                bid,
                fields["loan_number"],
                fields["loan_type"],
                fields["property_address"],
                fields["original_principal"],
                fields["current_balance"],
                fields["rate"],
                fields["pricing_mode"],
                fields["flat_fee"],
                fields["start_date"],
                fields["maturity_date"],
                fields["payment_type"],
                fields["payment_amount"],
                fields["payment_frequency"],
                fields["next_payment_due"],
                fields["status"],
                fields["base_term_months"],
            ),
        )
        brittco.db().commit()
        lid = brittco.db().execute(
            "SELECT id FROM loans WHERE loan_number=?", (fields["loan_number"],)
        ).fetchone()["id"]
    return bid, lid, fields


def _latest_schedule(bid):
    with brittco.app.app_context():
        doc = brittco.db().execute(
            """SELECT d.*, f.loan_id AS folder_loan_id
               FROM documents d
               LEFT JOIN doc_file_items i ON i.document_id = d.id
               LEFT JOIN doc_files f ON f.id = i.file_id
               WHERE d.borrower_id=? AND d.kind='Amortization schedule'
               ORDER BY d.id DESC LIMIT 1""",
            (bid,),
        ).fetchone()
        assert doc is not None
        info = {key: doc[key] for key in doc.keys()}
    path = os.path.join(brittco.UPLOAD_DIR, info["filename"])
    assert os.path.isfile(path)
    return info, path


def _sheet(path):
    wb = load_workbook(path)
    ws = wb.active
    labels = {}
    header_row = None
    for row in ws.iter_rows(min_row=1, max_col=8):
        label = row[0].value
        if label == "Payment #":
            header_row = row[0].row
            break
        if isinstance(label, str):
            labels[label] = row[1].value
    payments = []
    for row in ws.iter_rows(min_row=header_row + 1, max_col=8, values_only=True):
        if row[0] is None:
            break
        payments.append(row)
    return labels, payments


def test_schedule_math_uses_loan_terms():
    rows = build_amortization_schedule(
        {
            "original_principal": 120000,
            "current_balance": 120000,
            "rate": 12,
            "pricing_mode": "rate",
            "payment_amount": 1200,
            "payment_type": "Interest only",
            "loan_type": "Fix and Flip",
            "payment_frequency": "Monthly",
            "start_date": "2026-01-15",
            "maturity_date": "2026-04-15",
            "next_payment_due": "2026-02-15",
            "base_term_months": 3,
        }
    )
    assert [row["due_date"] for row in rows] == ["2026-02-15", "2026-03-15", "2026-04-15"]
    assert rows[0]["interest"] == 1200
    assert rows[0]["principal"] == 0
    assert rows[-1]["principal"] == 120000
    assert rows[-1]["ending_balance"] == 0


def test_flat_fee_is_due_at_maturity_not_as_a_rate():
    rows = build_amortization_schedule(
        {
            "original_principal": 50000,
            "current_balance": 50000,
            "rate": 99,
            "pricing_mode": "flat_fee",
            "flat_fee": 3000,
            "payment_amount": 0,
            "payment_type": "Interest only",
            "loan_type": "Fix and Flip",
            "payment_frequency": "Interest due at maturity",
            "start_date": "2026-01-15",
            "maturity_date": "2026-04-15",
            "next_payment_due": "2026-04-15",
        }
    )
    assert len(rows) == 1
    assert rows[0]["due_date"] == "2026-04-15"
    assert rows[0]["interest"] == 3000
    assert rows[0]["principal"] == 50000


def test_document_name_includes_loan_borrower_and_date():
    name = amortization_document_name(
        {"id": 42, "loan_number": "BC-100"},
        "Ada Borrower",
    )
    assert name.startswith("Amortization-42-BC-100-Ada-Borrower-")
    assert name.endswith(date.today().isoformat() + ".xlsx")


def test_button_is_on_the_staff_loan_page(monkeypatch):
    monkeypatch.setattr(brittco, "send_mail", lambda *a, **k: True)
    email = f"btn-{uuid.uuid4().hex}@example.com"
    _bid, lid, _fields = _insert_loan("Button Borrower", email)
    page = _staff().get(f"/loans/{lid}").get_data(as_text=True)
    assert ">Create amortization schedule<" in page
    assert f"/loans/{lid}/amortization-schedule" in page

    borrower = brittco.app.test_client()
    with borrower.session_transaction() as sess:
        sess["borrower_id"] = _bid
    portal = borrower.get("/portal/tools").get_data(as_text=True)
    assert "Create amortization schedule" not in portal
    investor = brittco.app.test_client()
    with investor.session_transaction() as sess:
        sess["investor_id"] = 1
    tools = investor.get("/investor/tools")
    if tools.status_code == 200:
        assert "Create amortization schedule" not in tools.get_data(as_text=True)


def test_create_files_schedule_and_emails_borrower(monkeypatch):
    sent = {}

    def fake_send(to_email, subject, body, attachment=None, attachment_name=None):
        sent["to"] = to_email
        sent["subject"] = subject
        sent["body"] = body
        sent["attachment"] = attachment
        sent["name"] = attachment_name
        return True

    monkeypatch.setattr(brittco, "send_mail", fake_send)
    email = f"ada-{uuid.uuid4().hex}@example.com"
    bid, lid, fields = _insert_loan("Ada Borrower", email)
    resp = _staff().post(f"/loans/{lid}/amortization-schedule", follow_redirects=True)
    page = resp.get_data(as_text=True)
    assert resp.status_code == 200
    assert "emailed to " + email in page

    info, path = _latest_schedule(bid)
    assert info["borrower_id"] == bid
    assert info["folder_loan_id"] == lid
    assert info["kind"] == "Amortization schedule"
    assert str(lid) in info["original_name"]
    assert "Ada-Borrower" in info["original_name"]
    assert date.today().isoformat() in info["original_name"]
    assert fields["loan_number"] in info["original_name"]

    labels, payments = _sheet(path)
    assert labels["Loan number"] == fields["loan_number"]
    assert labels["Borrower"] == "Ada Borrower"
    assert labels["Original principal"] == 120000
    assert labels["Rate %"] == 12
    assert payments[0][1] == "2026-02-15"
    assert payments[0][4] == 1200
    assert payments[-1][5] == 120000

    raw = open(path, "rb").read()
    assert sent["to"] == email
    assert sent["subject"] == f"Brittco Capital amortization schedule — {fields['loan_number']}"
    assert "amortization schedule" in sent["body"].lower()
    assert sent["attachment"] == raw
    assert sent["name"] == info["original_name"]


def test_missing_email_still_stores_the_schedule(monkeypatch):
    def fail_if_called(*_a, **_k):
        raise AssertionError("send_mail should not run without a borrower email")

    monkeypatch.setattr(brittco, "send_mail", fail_if_called)
    bid, lid, _fields = _insert_loan("No Mail", None)
    resp = _staff().post(f"/loans/{lid}/amortization-schedule", follow_redirects=True)
    page = resp.get_data(as_text=True)
    assert "No borrower email is on file" in page
    assert "saved to the borrower's documents" in page
    info, path = _latest_schedule(bid)
    assert os.path.getsize(path) > 100
    assert info["folder_loan_id"] == lid


def test_email_failure_still_stores_the_schedule(monkeypatch):
    monkeypatch.setattr(brittco, "send_mail", lambda *a, **k: False)
    email = f"fail-{uuid.uuid4().hex}@example.com"
    bid, lid, _fields = _insert_loan("Mail Fail", email)
    page = _staff().post(
        f"/loans/{lid}/amortization-schedule", follow_redirects=True
    ).get_data(as_text=True)
    assert "did not send" in page
    assert email in page
    info, _path = _latest_schedule(bid)
    assert info["borrower_id"] == bid


def test_borrower_cannot_create_schedule(monkeypatch):
    monkeypatch.setattr(brittco, "send_mail", lambda *a, **k: True)
    email = f"nope-{uuid.uuid4().hex}@example.com"
    bid, lid, _fields = _insert_loan("Portal User", email)
    client = brittco.app.test_client()
    with client.session_transaction() as sess:
        sess["borrower_id"] = bid
    resp = client.post(f"/loans/{lid}/amortization-schedule")
    assert resp.status_code in (302, 303)
    assert "login" in (resp.headers.get("Location") or "")
    anon = brittco.app.test_client().post(f"/loans/{lid}/amortization-schedule")
    assert anon.status_code in (302, 303)
    assert "login" in (anon.headers.get("Location") or "")
    with brittco.app.app_context():
        count = brittco.db().execute(
            "SELECT COUNT(*) AS n FROM documents WHERE borrower_id=? AND kind='Amortization schedule'",
            (bid,),
        ).fetchone()["n"]
    assert count == 0
