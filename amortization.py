"""Loan amortization schedule builder and XLSX export for Brittco Capital."""
from __future__ import annotations

import calendar
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _money(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _parse_date(s):
    if not s:
        return None
    if isinstance(s, date) and not isinstance(s, datetime):
        return s
    if isinstance(s, datetime):
        return s.date()
    try:
        return date.fromisoformat(str(s)[:10])
    except ValueError:
        return None


def _row_get(row, key, default=None):
    if row is None:
        return default
    try:
        if hasattr(row, "keys") and key in row.keys():
            v = row[key]
            return default if v is None else v
    except Exception:
        pass
    try:
        return row[key]
    except Exception:
        return default


def add_months(d, months):
    """Add calendar months, clamping day to end of month."""
    if not d:
        return None
    months = int(months)
    y = d.year + (d.month - 1 + months) // 12
    m = (d.month - 1 + months) % 12 + 1
    last = calendar.monthrange(y, m)[1]
    return date(y, m, min(d.day, last))


def _freq_step_months(freq):
    f = (freq or "Monthly").strip().lower()
    if "week" in f:
        return None  # weekly handled separately
    if "quarter" in f:
        return 3
    if "semi" in f or "6" in f:
        return 6
    if "year" in f or "annual" in f:
        return 12
    if "maturity" in f or "at payoff" in f or "due at" in f:
        return None  # single payment at maturity
    return 1  # monthly default


def _is_interest_only(loan):
    kind = (
        str(_row_get(loan, "payment_type") or "")
        + " "
        + str(_row_get(loan, "loan_type") or "")
    ).lower()
    if "amort" in kind:
        return False
    if "interest" in kind or "balloon" in kind or "standing" in kind or "transactional" in kind:
        return True
    # Default Brittco style is interest-only
    return True


def _monthly_interest(balance, rate_pct, payment_amount=None):
    if payment_amount and payment_amount > 0 and payment_amount < balance * 0.25:
        return round(_money(payment_amount), 2)
    rate = _money(rate_pct)
    return round(balance * rate / 100.0 / 12.0, 2)


def _amort_payment(principal, annual_rate_pct, n_periods):
    """Standard fixed monthly payment for amortizing loan."""
    p = _money(principal)
    n = int(n_periods or 0)
    if p <= 0 or n <= 0:
        return 0.0
    r = _money(annual_rate_pct) / 100.0 / 12.0
    if r <= 0:
        return round(p / n, 2)
    return round(p * (r * (1 + r) ** n) / ((1 + r) ** n - 1), 2)


def _period_dates(start, maturity, freq, payment_amount, next_due=None):
    """Generate due dates from start (or next_due) through maturity."""
    step = _freq_step_months(freq)
    if step is None and ("maturity" in (freq or "").lower() or "payoff" in (freq or "").lower() or "due at" in (freq or "").lower()):
        if maturity:
            return [maturity]
        return []
    if step is None:
        # weekly-ish: every 7 days
        dates = []
        cur = start or date.today()
        end = maturity or add_months(cur, 12)
        guard = 0
        while cur <= end and guard < 600:
            dates.append(cur)
            cur = cur + timedelta(days=7)
            guard += 1
        return dates

    first = next_due or (add_months(start, step) if start else date.today())
    if start and first < start:
        first = start
    end = maturity
    if not end:
        # standing / no maturity: project 12 more periods from today
        end = add_months(max(first, date.today()), step * 12)
    dates = []
    cur = first
    guard = 0
    while cur <= end and guard < 600:
        dates.append(cur)
        cur = add_months(cur, step)
        guard += 1
    if end and (not dates or dates[-1] < end):
        # ensure maturity included for balloon
        if not dates or dates[-1] != end:
            dates.append(end)
    return dates


def build_amortization_schedule(loan, payments=None):
    """
    Build schedule rows: history from payments, then remaining from current_balance
    through maturity.

    Each row: payment_no, due_date, beginning_balance, payment, interest, principal,
    ending_balance, status
    """
    payments = list(payments or [])
    payments = sorted(
        payments,
        key=lambda p: (_parse_date(_row_get(p, "paid_on")) or date.min, _row_get(p, "id") or 0),
    )

    original = _money(_row_get(loan, "original_principal"))
    current = _money(_row_get(loan, "current_balance"))
    if current <= 0 and original > 0:
        current = original
    rate = _money(_row_get(loan, "rate"))
    flat_fee = 0.0
    if str(_row_get(loan, "pricing_mode") or "").strip().lower() == "flat_fee":
        rate = 0.0
        flat_fee = _money(_row_get(loan, "flat_fee"))
    pay_amt = _money(_row_get(loan, "payment_amount"))
    pay_type = str(_row_get(loan, "payment_type") or "Interest only")
    freq = str(_row_get(loan, "payment_frequency") or "Monthly")
    start = _parse_date(_row_get(loan, "start_date"))
    maturity = _parse_date(_row_get(loan, "maturity_date"))
    next_due = _parse_date(_row_get(loan, "next_payment_due"))
    base_term = int(_money(_row_get(loan, "base_term_months")) or 0)
    interest_only = _is_interest_only(loan)

    rows = []
    bal = original if original > 0 else current
    n = 0

    # Historical payments
    for p in payments:
        paid_on = _parse_date(_row_get(p, "paid_on"))
        amt = _money(_row_get(p, "amount"))
        applied = str(_row_get(p, "applied_to") or "").strip()
        begin = bal
        if applied.lower() == "principal":
            interest = 0.0
            principal = min(amt, begin)
        elif applied.lower() == "interest":
            interest = amt
            principal = 0.0
        else:
            # Split: interest first, remainder principal
            interest = min(amt, _monthly_interest(begin, rate, pay_amt if interest_only else None))
            if interest_only and applied.lower() in ("", "other", "payment"):
                interest = min(amt, amt)  # typically all interest
                # if amount looks like IO payment, treat as interest
                expected_io = _monthly_interest(begin, rate, pay_amt)
                if abs(amt - expected_io) < 1.0 or amt <= expected_io * 1.05:
                    interest = amt
                    principal = 0.0
                else:
                    interest = min(amt, expected_io)
                    principal = max(0.0, round(amt - interest, 2))
            else:
                principal = max(0.0, round(amt - interest, 2))
        principal = min(principal, begin)
        end = round(begin - principal, 2)
        n += 1
        rows.append(
            {
                "payment_no": n,
                "due_date": paid_on.isoformat() if paid_on else "",
                "beginning_balance": round(begin, 2),
                "payment": round(amt, 2),
                "interest": round(interest, 2),
                "principal": round(principal, 2),
                "ending_balance": end,
                "status": "Paid",
            }
        )
        bal = end

    # Align remaining schedule start balance to current_balance when available
    if current > 0:
        bal = current

    today = date.today()
    # Remaining scheduled payments
    if bal <= 0.005:
        return rows

    # Build future due dates
    # Prefer next_payment_due; else first period after last paid / start
    last_paid = None
    if payments:
        last_paid = _parse_date(_row_get(payments[-1], "paid_on"))
    sched_start = next_due
    if not sched_start:
        if last_paid:
            step = _freq_step_months(freq) or 1
            sched_start = add_months(last_paid, step)
        elif start:
            step = _freq_step_months(freq) or 1
            sched_start = add_months(start, step)
        else:
            sched_start = today

    # Skip past due dates already covered by history (same date as paid)
    paid_dates = {
        (_parse_date(_row_get(p, "paid_on")) or date.min)
        for p in payments
    }

    due_dates = _period_dates(start, maturity, freq, pay_amt, next_due=sched_start)
    # Filter out dates already in paid history and dates before last paid
    filtered = []
    for d in due_dates:
        if d in paid_dates:
            continue
        if last_paid and d <= last_paid:
            continue
        filtered.append(d)
    due_dates = filtered

    # Amortizing payment amount
    if not interest_only:
        remaining_periods = len(due_dates) or base_term or 12
        if pay_amt > 0 and pay_amt < bal:
            amort_pmt = pay_amt
        else:
            amort_pmt = _amort_payment(bal, rate, remaining_periods)
    else:
        amort_pmt = pay_amt if pay_amt > 0 else _monthly_interest(bal, rate)

    for i, due in enumerate(due_dates):
        begin = bal
        is_last = (i == len(due_dates) - 1) or (maturity and due >= maturity)
        if interest_only:
            interest = _monthly_interest(begin, rate, pay_amt)
            if is_last and begin > 0.005:
                # balloon principal at maturity; flat fee is due then, not as a rate
                interest = round(interest + flat_fee, 2)
                principal = round(begin, 2)
                payment = round(interest + principal, 2)
            else:
                principal = 0.0
                payment = interest
        else:
            interest = round(begin * rate / 100.0 / 12.0, 2)
            if is_last and flat_fee:
                interest = round(interest + flat_fee, 2)
            payment = amort_pmt
            principal = round(payment - interest, 2)
            if principal > begin or is_last:
                principal = round(begin, 2)
                payment = round(interest + principal, 2)
            if principal < 0:
                principal = 0.0
                payment = interest
        end = round(max(0.0, begin - principal), 2)
        n += 1
        status = "Scheduled"
        if due < today:
            status = "Past due"
        rows.append(
            {
                "payment_no": n,
                "due_date": due.isoformat(),
                "beginning_balance": round(begin, 2),
                "payment": round(payment, 2),
                "interest": round(interest, 2),
                "principal": round(principal, 2),
                "ending_balance": end,
                "status": status,
            }
        )
        bal = end
        if bal <= 0.005:
            break

    return rows


def amortization_workbook(loan, borrower_name="", schedule=None, payments=None):
    """Return XLSX bytes for the amortization workbook."""
    if schedule is None:
        schedule = build_amortization_schedule(loan, payments=payments)

    wb = Workbook()
    ws = wb.active
    ws.title = "Amortization"

    header_font = Font(name="Calibri", bold=True, size=14, color="0D1520")
    label_font = Font(name="Calibri", bold=True, size=11)
    value_font = Font(name="Calibri", size=11)
    col_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    col_fill = PatternFill("solid", fgColor="0D1520")
    paid_fill = PatternFill("solid", fgColor="E8F5EE")
    thin = Border(
        left=Side(style="thin", color="D5DEE8"),
        right=Side(style="thin", color="D5DEE8"),
        top=Side(style="thin", color="D5DEE8"),
        bottom=Side(style="thin", color="D5DEE8"),
    )
    money_fmt = '#,##0.00'
    pct_fmt = '0.00"%"'

    loan_number = str(_row_get(loan, "loan_number") or f"Loan-{_row_get(loan, 'id')}")
    prop = str(_row_get(loan, "property_address") or "")
    original = _money(_row_get(loan, "original_principal"))
    current = _money(_row_get(loan, "current_balance"))
    flat = str(_row_get(loan, "pricing_mode") or "").strip().lower() == "flat_fee"
    rate = 0.0 if flat else _money(_row_get(loan, "rate"))
    flat_fee = _money(_row_get(loan, "flat_fee")) if flat else 0.0
    pay_amt = _money(_row_get(loan, "payment_amount"))
    pay_type = str(_row_get(loan, "payment_type") or "")
    freq = str(_row_get(loan, "payment_frequency") or "")
    start = str(_row_get(loan, "start_date") or "")
    maturity = str(_row_get(loan, "maturity_date") or "")
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")

    ws["A1"] = "Brittco Capital Inc"
    ws["A1"].font = header_font
    ws["A2"] = "Loan amortization schedule"
    ws["A2"].font = Font(name="Calibri", size=12, italic=True, color="5D6B7C")

    pricing_rows = [("Pricing", "Flat fee"), ("Flat fee", flat_fee)] if flat else [("Rate %", rate)]
    meta = [
        ("Loan number", loan_number),
        ("Borrower", borrower_name or ""),
        ("Property", prop),
        ("Original principal", original),
        ("Current balance", current),
        *pricing_rows,
        ("Payment amount", pay_amt),
        ("Payment type / frequency", f"{pay_type} / {freq}".strip(" /")),
        ("Start date", start[:10] if start else ""),
        ("Maturity date", maturity[:10] if maturity else ""),
        ("Generated at", generated),
    ]
    for i, (label, val) in enumerate(meta, start=4):
        ws.cell(row=i, column=1, value=label).font = label_font
        cell = ws.cell(row=i, column=2, value=val)
        cell.font = value_font
        if label in ("Original principal", "Current balance", "Payment amount", "Flat fee"):
            cell.number_format = money_fmt
        if label == "Rate %":
            cell.number_format = "0.00"

    headers = [
        "Payment #",
        "Due date",
        "Beginning balance",
        "Payment",
        "Interest",
        "Principal",
        "Ending balance",
        "Status",
    ]
    header_row = 4 + len(meta) + 1
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = col_font
        cell.fill = col_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin

    money_cols = {3, 4, 5, 6, 7}
    for r_i, row in enumerate(schedule, start=1):
        r = header_row + r_i
        vals = [
            row["payment_no"],
            row["due_date"],
            row["beginning_balance"],
            row["payment"],
            row["interest"],
            row["principal"],
            row["ending_balance"],
            row["status"],
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = thin
            cell.font = value_font
            if c in money_cols:
                cell.number_format = money_fmt
            if row["status"] == "Paid":
                cell.fill = paid_fill

    widths = [12, 14, 18, 14, 12, 12, 16, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A17"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), loan_number


def filename_for_loan(loan):
    num = str(_row_get(loan, "loan_number") or f"Loan-{_row_get(loan, 'id')}")
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in num).strip("-") or "loan"
    return f"Amortization-{safe}.xlsx"


def _safe_token(text, fallback=""):
    cleaned = "".join(ch if ch.isalnum() else "-" for ch in str(text or "").strip())
    while "--" in cleaned:
        cleaned = cleaned.replace("--", "-")
    return cleaned.strip("-") or fallback


def amortization_document_name(loan, borrower_name=""):
    """Borrower-file name: loan id, loan number, borrower, and the day it was created."""
    lid = _safe_token(_row_get(loan, "id"), "loan")
    number = _safe_token(_row_get(loan, "loan_number"), "")
    who = _safe_token(borrower_name, "borrower")
    bits = ["Amortization", lid]
    if number and number.lower() != lid.lower():
        bits.append(number)
    bits.extend([who, date.today().isoformat()])
    name = "-".join(bits) + ".xlsx"
    if len(name) > 180:
        name = name[:170].rstrip("-") + ".xlsx"
    return name
